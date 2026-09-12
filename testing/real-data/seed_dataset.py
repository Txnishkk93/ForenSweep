"""Seed a real test volume with synthetic, hashable forensic fixtures."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import random
import shutil
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

try:
	from .device_guard import DeviceInfo, require_device_gate
except ImportError:
	from device_guard import DeviceInfo, require_device_gate


SCRIPT_DIR = Path(__file__).resolve().parent
RESULTS_DIR = SCRIPT_DIR.parent / "results"
SCHEMA_PATH = SCRIPT_DIR / "dataset_schema.json"
MANIFEST_PATH = RESULTS_DIR / "seed_manifest.csv"


def sha256_file(path: Path) -> str:
	digest = hashlib.sha256()
	with path.open("rb") as handle:
		for block in iter(lambda: handle.read(1024 * 1024), b""):
			digest.update(block)
	return digest.hexdigest()


def _write_text(path: Path, size_kb: int, pii: bool = False) -> None:
	if pii:
		text = "Synthetic PII fixture\nName: Ada Example\nEmail: ada@example.invalid\nPhone: 000-000-0000\n"
	else:
		text = "ForenSweep synthetic forensic fixture.\n"
	encoded = text.encode()
	path.write_bytes((encoded * ((size_kb * 1024 // len(encoded)) + 1))[: size_kb * 1024])


def _write_pdf(path: Path, size_kb: int) -> None:
	from reportlab.pdfgen.canvas import Canvas

	canvas = Canvas(str(path))
	canvas.drawString(72, 720, "ForenSweep synthetic PDF fixture")
	canvas.save()
	if path.stat().st_size < size_kb * 1024:
		with path.open("ab") as handle:
			handle.write(b"\n% synthetic padding\n" * ((size_kb * 1024 - path.stat().st_size) // 20 + 1))


def _write_docx(path: Path, size_kb: int) -> None:
	from docx import Document

	document = Document()
	document.add_heading("ForenSweep synthetic document", level=1)
	document.add_paragraph("Synthetic evidence fixture; not real personal data.")
	document.save(path)


def _write_xlsx(path: Path, size_kb: int) -> None:
	from openpyxl import Workbook

	workbook = Workbook()
	sheet = workbook.active
	sheet["A1"] = "ForenSweep synthetic spreadsheet"
	sheet["A2"] = "Synthetic evidence fixture"
	workbook.save(path)


def _write_image(path: Path, size_kb: int) -> None:
	from PIL import Image, ImageDraw

	image = Image.new("RGB", (320, 200), (220, 225, 230))
	ImageDraw.Draw(image).text((20, 90), "ForenSweep", fill=(20, 25, 30))
	image.save(path, format="JPEG" if path.suffix.lower() == ".jpg" else "PNG")


def generate_file(path: Path, extension: str, size_kb: int, synthetic_pii: bool = False) -> str:
	path.parent.mkdir(parents=True, exist_ok=True)
	ext = extension.lower().lstrip(".")
	if ext == "pdf":
		_write_pdf(path, size_kb)
	elif ext == "docx":
		_write_docx(path, size_kb)
	elif ext == "xlsx":
		_write_xlsx(path, size_kb)
	elif ext in {"jpg", "png"}:
		_write_image(path, size_kb)
	elif ext == "zip":
		with ZipFile(path, "w", ZIP_DEFLATED) as archive:
			archive.writestr("README.txt", "ForenSweep synthetic archive fixture\n")
	elif ext == "mp4":
		path.write_bytes(b"ForenSweep synthetic video fixture\n" + os.urandom(min(size_kb * 1024, 1024)))
	else:
		_write_text(path, size_kb, synthetic_pii)
	# Hash immediately, before any delete pass can change the filesystem.
	return sha256_file(path)


def _manifest_rows(schema: dict[str, Any], root: Path) -> list[dict[str, str]]:
	rows: list[dict[str, str]] = []
	rng = random.Random(42)
	for folder in schema["folders"]:
		low, high = folder["size_range_kb"]
		for index in range(folder["files"]):
			extension = folder["types"][index % len(folder["types"])]
			size_kb = rng.randint(low, high)
			relative = Path(folder["path"]) / f"fixture_{index + 1:03d}.{extension}"
			path = root / relative
			file_hash = generate_file(path, extension, size_kb, folder.get("synthetic_pii", False))
			rows.append({"path": str(relative), "sha256": file_hash, "deleted": "false", "size_bytes": str(path.stat().st_size)})
	return rows


def write_manifest(rows: list[dict[str, str]]) -> None:
	RESULTS_DIR.mkdir(parents=True, exist_ok=True)
	with MANIFEST_PATH.open("w", newline="", encoding="utf-8") as handle:
		writer = csv.DictWriter(handle, fieldnames=["path", "sha256", "deleted", "size_bytes"])
		writer.writeheader()
		writer.writerows(rows)


def apply_deletes(
	root: Path,
	patterns: list[dict[str, Any]],
	rows: list[dict[str, str]] | None = None,
) -> None:
	"""Delete selected files and only flip existing manifest rows to deleted=true."""
	if rows is None:
		with MANIFEST_PATH.open(newline="", encoding="utf-8") as handle:
			rows = list(csv.DictReader(handle))
	for rule in patterns:
		candidates = [row for row in rows if Path(row["path"]).parent == Path(rule["folder"]) and Path(row["path"]).match(rule["pattern"]) and row["deleted"] == "false"]
		count = int(len(candidates) * float(rule["percent"]) / 100)
		for row in sorted(candidates, key=lambda item: item["path"])[:count]:
			target = root / row["path"]
			if target.exists():
				target.unlink()
			row["deleted"] = "true"
	write_manifest(rows)


def check_budget(schema: dict[str, Any], device: DeviceInfo, root: Path) -> None:
	projected = sum(folder["files"] * folder["size_range_kb"][1] * 1024 for folder in schema["folders"])
	usage = shutil.disk_usage(root)
	contributors = sorted(((folder["path"], folder["files"] * folder["size_range_kb"][1] * 1024) for folder in schema["folders"]), key=lambda item: item[1], reverse=True)
	print(f"Budget: projected={projected / 1024**3:.2f} GiB, available={usage.free / 1024**3:.2f} GiB, device={device.path}")
	print("Largest schema contributors:", ", ".join(f"{name}={amount / 1024**3:.2f} GiB" for name, amount in contributors[:3]))
	if projected > usage.free * 0.8:
		raise RuntimeError("Seed budget exceeds 80% of actual free space; refusing to write the volume")


def main() -> int:
	parser = argparse.ArgumentParser()
	parser.add_argument("--volume", required=True, help="Mounted test volume path; a folder is not enough for device identity")
	parser.add_argument("--i-am-sure", dest="expected_serial", help="Exact physical-device serial")
	args = parser.parse_args()
	# CRITICAL: this resolves the physical device and checks size/system-disk identity before any write/delete.
	device = require_device_gate(args.volume, args.expected_serial, RESULTS_DIR)
	schema = json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))
	root = Path(args.volume).resolve()
	check_budget(schema, device, root)
	rows = _manifest_rows(schema, root)
	# Finalize every pre-delete hash and row before applying any deletion rule.
	write_manifest(rows)
	assert rows and all(row["sha256"] and row["deleted"] in {"false", "true"} for row in rows)
	apply_deletes(root, schema.get("delete_patterns", []), rows)
	final_rows = list(csv.DictReader(MANIFEST_PATH.open(encoding="utf-8")))
	assert all(row["sha256"] for row in final_rows), "Every manifest row must retain its pre-delete SHA-256"
	print(f"Seeded {len(final_rows)} files; manifest: {MANIFEST_PATH}")
	return 0


if __name__ == "__main__":
	raise SystemExit(main())
